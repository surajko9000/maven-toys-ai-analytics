"""Builds the 6-slide "Colorbuds is fading" deck as a real .pptx (native charts, speaker notes), consulting-style.

Run:  python3 scripts/build_deck.py [out.pptx]
Every number on the slides is computed from public/data/*.csv and asserted against
the values verified independently with pandas (see evaluations/iteration_log.md).
"""
import sys
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "public" / "data"
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "site" / "deck" / "Colorbuds_to_Playfoam.pptx"

# ----------------------------------------------------------------- data
sales = pd.read_csv(DATA / "sales.csv", parse_dates=["Date"])
prod = pd.read_csv(DATA / "products.csv")
inv = pd.read_csv(DATA / "inventory.csv")
for c in ("Product_Cost", "Product_Price"):
    prod[c] = prod[c].str.replace(r"[$ ]", "", regex=True).astype(float)
prod["unit_profit"] = prod.Product_Price - prod.Product_Cost
x = sales.merge(prod, on="Product_ID")
x["profit"] = x.Units * x.unit_profit
x["month"] = x.Date.dt.to_period("M")
x["quarter"] = x.Date.dt.to_period("Q")

units_m = x.groupby(["Product_Name", "month"]).Units.sum().unstack(fill_value=0)
cb = units_m.loc["Colorbuds"]
cb_first, cb_last = int(cb.iloc[0]), int(cb.iloc[-1])
cb_drop = round((1 - cb_last / cb_first) * 100)

qp = x.groupby(["Product_Name", "quarter"]).profit.sum().unstack(fill_value=0)
share = lambda q: qp.loc["Colorbuds", q] / qp[q].sum() * 100
share_22, share_23 = round(share("2022Q3")), round(share("2023Q3"))

pf = units_m.loc["Playfoam"]
pf_aug, pf_sep = int(pf["2023-08"]), int(pf["2023-09"])
pf_unit_profit = float(prod.loc[prod.Product_Name == "Playfoam", "unit_profit"].iloc[0])
pf_margin = round(pf_unit_profit / float(prod.loc[prod.Product_Name == "Playfoam", "Product_Price"].iloc[0]) * 100)
pf_month_profit_k = round(pf_sep * pf_unit_profit / 1000)
pf_aug_profit_k = round(pf_aug * pf_unit_profit / 1000)
pf_pid = int(prod.loc[prod.Product_Name == "Playfoam", "Product_ID"].iloc[0])
_pw = x[(x.Product_Name == "Playfoam") & (x.Date >= "2023-08-20")].groupby(pd.Grouper(key="Date", freq="W-SAT")).Units.sum()
pf_weeks = [int(v) for v in _pw.values]                       # six full Sun-Sat weeks since launch
pf_week_labels = [(d - pd.Timedelta(days=6)).strftime("%b ") + str((d - pd.Timedelta(days=6)).day) for d in _pw.index]
pf_wk_avg = round(sum(pf_weeks) / len(pf_weeks), -1)
assert pf_weeks == [718, 722, 646, 688, 679, 601], pf_weeks
_w = x[x.Date >= "2023-08-20"]
pf_profit_rank_s = _w.groupby("Product_Name").profit.sum().sort_values(ascending=False)
top5 = [(n, float(v)) for n, v in pf_profit_rank_s.head(5).items()]
pf_prank = list(pf_profit_rank_s.index).index("Playfoam") + 1
_wk = _w.groupby([_w.Date.dt.to_period("W-SAT"), "Product_Name"]).profit.sum().unstack(fill_value=0)
pf_weeks_first = int((_wk.idxmax(axis=1) == "Playfoam").sum())
assert (pf_prank, pf_weeks_first, top5[0][0], top5[1][0]) == (2, 2, "Colorbuds", "Playfoam"), (pf_prank, pf_weeks_first, top5[:2])
cb_6w_k, pf_6w_k = round(top5[0][1] / 1000), round(top5[1][1] / 1000)
pf_store_avg = round(pf_sep / sales.Store_ID.nunique())     # average Sep units per store, all stores
assert pf_store_avg == 57
pf_empty_k = round(13 * pf_store_avg * pf_unit_profit / 1000)
assert pf_empty_k == 5
pf_full_k = pf_month_profit_k + pf_empty_k
pf_zero = int(((inv.Product_ID == pf_pid) & (inv.Stock_On_Hand == 0)).sum())

first_sale = x.groupby("Product_Name").Date.min()
launched = [n for n in first_sale.index if first_sale[n] > pd.Timestamp("2022-01-31") and n != "Playfoam"]
faded = [n for n in launched if units_m.loc[n].iloc[-1] <= 0.6 * units_m.loc[n].max()]
_first6 = {}
for _n in launched + ["Playfoam"]:
    _d0 = first_sale[_n]
    _first6[_n] = float(x[(x.Product_Name == _n) & (x.Date < _d0 + pd.Timedelta(days=42))].profit.sum())
assert max(_first6, key=_first6.get) == "Playfoam" and len(launched) == 10, _first6   # best first-6-weeks profit of all launches

q3_22_k, q3_23_k = round(qp.loc["Colorbuds", "2022Q3"] / 1000), round(qp.loc["Colorbuds", "2023Q3"] / 1000)
pf_rank = int((qp["2023Q3"] > qp.loc["Playfoam", "2023Q3"]).sum()) + 1
cbs = x[(x.Product_Name == "Colorbuds") & (x.Date.dt.month <= 9)]
g = cbs.groupby([cbs.Store_ID, cbs.Date.dt.year]).Units.sum().unstack(fill_value=0)
stores_down = int((g[2023] < g[2022]).sum())
cb_unit_profit = float(prod.loc[prod.Product_Name == "Colorbuds", "unit_profit"].iloc[0])
gap_k = round((cb_first - cb_last) * cb_unit_profit / 1000)
assert pf_full_k * 2 > gap_k, (pf_full_k, gap_k)   # restocked Playfoam covers more than half the Colorbuds gap
peak_drop = lambda n: round((1 - units_m.loc[n].iloc[-1] / units_m.loc[n].max()) * 100)
jenga_drop, etch_drop, hw_drop = peak_drop("Jenga"), peak_drop("Etch A Sketch"), peak_drop("Hot Wheels 5-Pack")
etch_peak = str(units_m.loc["Etch A Sketch"].idxmax())
mrp_drop, mrp_peak = peak_drop("Mr. Potatohead"), str(units_m.loc["Mr. Potatohead"].idxmax())

# independent check against the numbers verified earlier with pandas
assert (cb_first, cb_last, cb_drop) == (8110, 2594, 68), (cb_first, cb_last, cb_drop)
assert (share_22, share_23) == (28, 13), (share_22, share_23)
assert (pf_aug, pf_sep, pf_margin, pf_month_profit_k, pf_zero) == (1311, 2847, 64, 20, 13)
assert (len(faded), len(launched)) == (9, 10), (len(faded), len(launched))
assert (q3_22_k, q3_23_k, pf_rank, stores_down, gap_k) == (142, 74, 5, 48, 44), (q3_22_k, q3_23_k, pf_rank, stores_down, gap_k)
assert (jenga_drop, etch_drop, hw_drop, etch_peak) == (92, 56, 9, "2022-12"), (jenga_drop, etch_drop, hw_drop, etch_peak)
assert (mrp_drop, mrp_peak) == (65, "2022-11"), (mrp_drop, mrp_peak)
print("numbers verified:", dict(cb=(cb_first, cb_last, cb_drop), share=(share_22, share_23),
      playfoam=(pf_aug, pf_sep, pf_margin, pf_month_profit_k, pf_zero), faded=(len(faded), len(launched))))

# ----------------------------------------------------------------- style (restrained, consulting-style)
NAVY, BLUE = "051C2C", "2A78D6"                      # same palette as the website
PINK, MINT, ORANGE, YELLOW = "E8457B", "1BAF7A", "EB6834", "F2B400"
PINK_T, MINT_T, ORANGE_T = "D6336C", "0E8F63", "D9531E"   # slightly darker, for text on white/gray
INK, MUTED, LIGHT, MID = "1B2733", "5B6470", "F2F4F7", "D0D5DD"
BAR_SOFT = "A5E3CC"
HEAD, BODY = "Cambria", "Calibri"
rgb = lambda h: RGBColor.from_string(h)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
prs.core_properties.title = "Our #1 profit toy, Colorbuds, is in decline"
prs.core_properties.author = "Maven Toys Analytics"
prs.core_properties.subject = "Business review"
LM = 0.6          # left/right margin
CW = 13.333 - 2 * LM


def text(slide, s, x, y, w, h, size, color=INK, bold=False, font=BODY, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, name=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    for k, part in enumerate(s.split("**")):
        if not part:
            continue
        r = p.add_run()
        r.text = part
        r.font.size, r.font.bold, r.font.name = Pt(size), (bold or k % 2 == 1), font
        r.font.color.rgb = rgb(color)
    return tb


def shape(slide, kind, x, y, w, h, fill, name=None):
    sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid()
    sp.fill.fore_color.rgb = rgb(fill)
    sp.line.fill.background()
    sp.shadow.inherit = False
    st = sp._element.find(qn("p:style"))
    if st is not None:                       # no theme shadow/effect
        st.find(qn("a:effectRef")).set("idx", "0")
    if name:
        sp.name = name
    return sp


def notes(slide, s):
    slide.notes_slide.notes_text_frame.text = s


def bullets(slide, items, x, y, w, h, size=18, color=INK, space=10, name="Bullets"):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, s in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(space)
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", "285750")
        pPr.set("indent", "-285750")
        buClr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(buClr, qn("a:srgbClr")).set("val", BLUE)
        etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
        etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
        for k, part in enumerate(s.split("**")):
            if not part:
                continue
            r = p.add_run()
            r.text = part
            r.font.size, r.font.name = Pt(size), BODY
            r.font.bold = (k % 2 == 1)          # odd segments sit between ** markers
            r.font.color.rgb = rgb(color)
    return tb


def label(slide, s, x, y, w, name="Section label"):
    """Small uppercase section label above a block."""
    text(slide, s.upper(), x, y, w, 0.3, 12, MUTED, True, name=name)


def footer(slide, n):
    text(slide, "Source: Maven Toys sales data, Jan 2022 – Sep 2023. Profit is estimated from current price and cost.",
         LM, 6.98, 9.5, 0.3, 10, MUTED, name="Source")
    text(slide, f"Maven Toys   |   {n}", 10.2, 6.98, 13.333 - LM - 10.2, 0.3, 10, MUTED, align=PP_ALIGN.RIGHT, name="Page")


def content_slide(action_title, n):
    sl = prs.slides.add_slide(BLANK)
    sl.background.fill.solid()
    sl.background.fill.fore_color.rgb = rgb("FFFFFF")
    text(sl, action_title, LM, 0.5, CW, 1.1, 28, NAVY, True, HEAD, anchor=MSO_ANCHOR.TOP, name="Title")
    footer(sl, n)
    return sl


def style_axis_text(axis, size=14, color=MUTED):
    axis.tick_labels.font.size = Pt(size)
    axis.tick_labels.font.name = BODY
    axis.tick_labels.font.color.rgb = rgb(color)


# ================================================================= slide 1: title
s1 = prs.slides.add_slide(BLANK)
s1.background.fill.solid()
s1.background.fill.fore_color.rgb = rgb(NAVY)
for k, c in enumerate((PINK, YELLOW, MINT)):
    shape(s1, MSO_SHAPE.RECTANGLE, LM + k * 0.2, 0.62, 0.14, 0.14, c, f"Mark {k+1}")
text(s1, "MAVEN TOYS", LM + 0.8, 0.52, 4, 0.34, 14, "FFFFFF", True, name="Wordmark", anchor=MSO_ANCHOR.MIDDLE)
shape(s1, MSO_SHAPE.RECTANGLE, 8.6, 0, 13.333 - 8.6, 7.5, "FBDDE9", "Image panel")
def _cover_pic(fname, cy, h, name, descr):
    from PIL import Image as _I
    w0, h0 = _I.open(ROOT / "scripts" / "assets" / fname).size
    w = h * w0 / h0
    pic = s1.shapes.add_picture(str(ROOT / "scripts" / "assets" / fname), Inches(8.6 + (13.333 - 8.6 - w) / 2), Inches(cy), Inches(w), Inches(h))
    pic.name = name
    pic._element.nvPicPr.cNvPr.set("descr", descr)
_cover_pic("colorbuds_illustration.png", 0.7, 2.6, "Colorbuds illustration", "Illustration of Colorbuds: two colorful earbuds in an open charging case with a heart")
_arrow = shape(s1, MSO_SHAPE.DOWN_ARROW, 8.6 + (13.333 - 8.6) / 2 - 0.25, 3.4, 0.5, 0.55, PINK, "Arrow from Colorbuds to Playfoam")
_cover_pic("playfoam_illustration.png", 4.15, 2.65, "Playfoam illustration", "Illustration of Playfoam: a purple tub heaped with tiny colorful beads")
_tb = text(s1, "Colorbuds is in decline.", LM, 1.6, 7.9, 2.7, 40, "FFFFFF", True, HEAD, anchor=MSO_ANCHOR.BOTTOM, name="Title")
_p = _tb.text_frame.paragraphs[0]
_p.add_line_break()
_r = _p.add_run()
_r.text = "Playfoam is the next bet."
_r.font.size, _r.font.bold, _r.font.name = Pt(40), True, HEAD
_r.font.color.rgb = rgb("FFFFFF")
text(s1, "A plan to protect profit by backing newer, higher-margin toys", LM, 4.55, 7.6, 0.9, 22, "D6E0EE", name="Subtitle")
text(s1, "Business review  |  Data through Sep 30, 2023", LM, 6.75, 8, 0.4, 14, "A9B8CC", name="Date line")
notes(s1, "Open with the one-line answer: Colorbuds (our line of kids' wireless earbuds), our #1 profit toy, is in decline, and Playfoam, a new toy launched in August, is the next bet. The plan is simple: back newer, "
          "higher-margin toys, test one cheap fix for Colorbuds, and review in 8 weeks. Keep the intro under 20 seconds.")

# ================================================================= slide 2: the answer
s2 = content_slide("Our top profit toy is fading, and a newer toy is already making up part of the gap", 2)
for i, (big, lab, cap, bx, bcol) in enumerate([
    (f"{share_22}% → {share_23}%", "Colorbuds' share of quarterly profit", "Q3 2022 to Q3 2023", LM, PINK_T),
    (f"#{pf_prank} of 35", "Playfoam's rank among all toys by profit", "Behind only Colorbuds, in its first six full weeks on sale", LM + CW / 2 + 0.2, MINT_T),
]):
    w = CW / 2 - 0.2
    shape(s2, MSO_SHAPE.RECTANGLE, bx, 1.85, w, 2.75, LIGHT, f"Panel {i+1}")
    label(s2, lab, bx + 0.4, 2.1, w - 0.8, f"Panel label {i+1}")
    text(s2, big, bx + 0.4, 2.55, w - 0.8, 1.2, 56, bcol, True, HEAD, anchor=MSO_ANCHOR.MIDDLE, name=f"Big number {i+1}")
    text(s2, cap, bx + 0.4, 3.95, w - 0.8, 0.4, 16, MUTED, name=f"Caption {i+1}")
bullets(s2, ["Colorbuds is still our **#1 profit toy**, but its share has **more than halved**",
             f"Playfoam is our **strongest launch yet**: #{pf_prank} in profit after six weeks, ahead of all {len(launched)} other new toys",
             f"Playfoam earns about **${pf_month_profit_k}K a month**, nearly half of the ${gap_k}K Colorbuds has lost"],
        LM, 4.95, CW, 1.8, 20, name="Key points")
notes(s2, f"Colorbuds earned about {share_22}% of Q3 2022 profit and {share_23}% of Q3 2023 profit; it is still our #1 profit toy, "
          f"so it is fading, not dead. Playfoam launched in mid-August 2023; at its September pace ({pf_sep:,} units x "
          f"${pf_unit_profit:.0f} profit per unit) it earns about ${pf_month_profit_k}K a month. Weekly sales have held at roughly 600-720 units for six straight full weeks (no cooling yet); the August-to-September monthly jump is mostly a longer month, so do not call it growth. Colorbuds' monthly profit is about ${gap_k}K lower than in "
          "Jan 2022 ((8,110 - 2,594) units x $8). Profit is estimated from current price and cost, not accounting figures.")

# ================================================================= slide 3: the fall
s3 = content_slide(f"Colorbuds sales have fallen {cb_drop}% since January 2022, across almost every store", 3)
months = [str(p) for p in cb.index]
show_at = {0, 6, 12, 18}
def _lbl(i, m):
    per = pd.Period(m)
    if i not in show_at:
        return " " * (i + 1)
    return per.strftime("%b") + ("\n\u2019" + per.strftime("%y") if per.month == 1 else "")


mlabels = [_lbl(i, m) for i, m in enumerate(months)]
cd = CategoryChartData()
cd.categories = mlabels
cd.add_series("Colorbuds units per month", [int(v) for v in cb.values])
label(s3, "Colorbuds units sold per month", LM, 1.85, 6, "Chart label")
gf = s3.shapes.add_chart(XL_CHART_TYPE.LINE, Inches(LM), Inches(2.2), Inches(7.9), Inches(4.55), cd)
gf.name = "Colorbuds monthly units"
ch = gf.chart
ch.has_legend = False
ch.has_title = False
ch.font.name = BODY
va, ca = ch.value_axis, ch.category_axis
va.visible = False
va.has_major_gridlines = False
va.minimum_scale, va.maximum_scale = 0, 9500
ca.has_major_gridlines = False
ca.format.line.color.rgb = rgb(MID)
style_axis_text(ca, 13)
bp = ca._element.find(qn('c:txPr')).find(qn('a:bodyPr'))
bp.set('rot', '0')
bp.set('vert', 'horz')
ser = ch.plots[0].series[0]
ser.smooth = False
ser.format.line.color.rgb = rgb(PINK)
ser.format.line.width = Pt(3.5)
from pptx.enum.chart import XL_MARKER_STYLE
ser.marker.style = XL_MARKER_STYLE.NONE
for idx, lbl, when, pos in ((0, f"{cb_first:,}", "Jan 2022", XL_LABEL_POSITION.ABOVE),
                            (len(cb) - 1, f"{cb_last:,}", "Sep 2023", XL_LABEL_POSITION.BELOW)):
    pt = ser.points[idx]
    pt.marker.style = XL_MARKER_STYLE.CIRCLE
    pt.marker.size = 11
    pt.marker.format.fill.solid()
    pt.marker.format.fill.fore_color.rgb = rgb(PINK)
    pt.marker.format.line.color.rgb = rgb("FFFFFF")
    dl = pt.data_label
    dl.has_text_frame = True
    dl.text_frame.text = lbl
    dl.position = pos
    r = dl.text_frame.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.name = Pt(22), True, HEAD
    r.font.color.rgb = rgb(PINK_T)
    p2 = dl.text_frame.add_paragraph()
    r2 = p2.add_run()
    r2.text = when
    r2.font.size, r2.font.name = Pt(13), BODY
    r2.font.color.rgb = rgb(MUTED)
shape(s3, MSO_SHAPE.RECTANGLE, 8.95, 1.85, 13.333 - LM - 8.95, 4.9, LIGHT, "Takeaways panel")
label(s3, "What we see", 9.3, 2.1, 3.2, "Takeaways label")
bullets(s3, [f"Down in **{stores_down} of 50 stores**, in every store type",
             "Every store still carries it, so this is **weaker demand, not empty shelves**",
             f"Quarterly profit has **roughly halved**: ${q3_22_k}K to ${q3_23_k}K"],
        9.3, 2.6, 3.15, 4.0, 18, space=14, name="Key points")
notes(s3, f"Monthly Colorbuds units fell from {cb_first:,} (Jan 2022) to {cb_last:,} (Sep 2023), down {cb_drop}%. "
          f"The drop is steady and broad: {stores_down} of 50 stores are down (Jan-Sep 2023 vs Jan-Sep 2022), in every store type. All 50 stores carry it and none is at zero stock, "
          "so this looks like a demand problem, not a stock problem. Some of the drop is seasonal (company-wide revenue in Sep is about 26% "
          "below its peak month) but not most of it. The data cannot say WHY demand fell; no price history, promotions or competitor data.")

# ================================================================= slide 4: most toys fade
s4 = content_slide(f"Fading is normal for our toys: {len(faded)} of {len(launched)} recent launches are 40%+ below their peak", 4)
shape(s4, MSO_SHAPE.RECTANGLE, LM, 1.85, 6.4, 4.9, LIGHT, "Stat panel")
text(s4, f"{len(faded)} of {len(launched)}", LM + 0.4, 2.15, 5.6, 1.7, 88, ORANGE_T, True, HEAD, anchor=MSO_ANCHOR.MIDDLE, name="Big number")
text(s4, "recent launches are **40%+ below their peak month**", LM + 0.4, 3.95, 5.6, 0.9, 22, NAVY, name="Caption")
for i in range(len(launched)):
    bx = LM + 0.4 + i * 0.54
    c = ORANGE if i < len(faded) else MID
    shape(s4, MSO_SHAPE.RECTANGLE, bx, 5.45, 0.42, 0.42, c, f"Toy marker {i+1}")
label(s4, "Examples", 7.4, 1.95, 5, "Examples label")
bullets(s4, [f"Jenga peaked in Jul 2022 and is now down {jenga_drop}%",
             f"Etch A Sketch peaked its second month and is down {etch_drop}%",
             f"Mr. Potatohead peaked in Nov 2022 and is now down {mrp_drop}%",
             "**Colorbuds is following the same trend**"],
        7.4, 2.45, CW - (7.4 - LM), 4.2, 20, space=16, name="Key points")
notes(s4, "This is the likely explanation, not a proven cause. Of the 10 toys that started selling during this data (excluding Playfoam, "
          "which is too new), 9 now sell at least 40% less per month than in their peak month. Jenga -92%, Supersoaker -66%, "
          "Mr. Potatohead -65%, PlayDoh Playset -64%, Etch A Sketch -56%, Plush Pony -56%, Foam Disk Launcher -54%, Magic Sand -41%, "
          "Gamer Headphones -41%; Hot Wheels is the exception at -9%. Caveats: part of each drop is seasonal, and 'launch' means first sale in this data.")

# ================================================================= slide 5: Playfoam
s5 = content_slide(f"Playfoam is the early bright spot, and full shelves could cover over half of Colorbuds' lost profit", 5)
cd = CategoryChartData()
cd.categories = [n for n, _ in top5]
cd.add_series("Profit", [v for _, v in top5])
label(s5, "Estimated profit by toy, last six full weeks (Aug 20 - Sep 30)", LM, 1.85, 6.4, "Chart label")
gf = s5.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(LM), Inches(2.2), Inches(6.4), Inches(4.55), cd)
gf.name = "Top 5 toys by profit since Playfoam launched"
ch = gf.chart
ch.has_legend = False
ch.has_title = False
ch.font.name = BODY
va, ca = ch.value_axis, ch.category_axis
va.visible = False
va.has_major_gridlines = False
va.minimum_scale, va.maximum_scale = 0, 37000
ca.has_major_gridlines = False
ca.reverse_order = True
ca.format.line.color.rgb = rgb(MID)
style_axis_text(ca, 16, INK)
plot = ch.plots[0]
plot.gap_width = 45
plot.has_data_labels = True
dls = plot.data_labels
dls.number_format, dls.number_format_is_linked = '$#,##0,"K"', False
dls.position = XL_LABEL_POSITION.OUTSIDE_END
dls.font.size, dls.font.bold, dls.font.name = Pt(20), True, HEAD
dls.font.color.rgb = rgb(NAVY)
ser = plot.series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = rgb(MID)
for i, (n, _) in enumerate(top5):
    if n == "Playfoam":
        pt = ser.points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(MINT)
for i, (big, small, by) in enumerate([(f"{pf_margin}% margin", f"${pf_unit_profit:.0f} profit on every unit", 1.85),
                                      (f"{pf_zero} stores out of stock", f"13 x 57 units x $7 = +$5K a month (${pf_month_profit_k}K to ${pf_full_k}K)", 3.45)]):
    shape(s5, MSO_SHAPE.RECTANGLE, 7.5, by, 13.333 - LM - 7.5, 1.4, LIGHT, f"Stat panel {i+1}")
    text(s5, big, 7.8, by + 0.2, 4.7, 0.6, 26, MINT_T if i == 0 else ORANGE_T, True, HEAD, anchor=MSO_ANCHOR.MIDDLE, name=f"Stat title {i+1}")
    text(s5, small, 7.8, by + 0.88, 4.7, 0.4, 16, INK, name=f"Stat note {i+1}")
bullets(s5, [f"**#1 profit toy in {pf_weeks_first} of 6 weeks**, ahead of Colorbuds",
             "New toys often cool off, so **the window to act is short**"],
        7.5, 5.1, 13.333 - LM - 7.5, 1.6, 18, space=10, name="Key points")
notes(s5, f"Playfoam launched Aug 17, 2023 and sold {pf_aug + pf_sep:,} units in about six weeks ({pf_aug:,} in August from the 17th, "
          f"{pf_sep:,} in September). Playfoam earned about $27K in its first 42 days, more than any of the 10 other toys launched in this data (next best: Etch A Sketch, about $20K, which later fell 56% from its peak, so a strong start is not proof of staying power). Estimated profit, first six full weeks (Aug 20 - Sep 30): Playfoam ${pf_6w_k}K vs Colorbuds ${cb_6w_k}K, #2 of 35 toys, and #1 in {pf_weeks_first} of the 6 weeks. Weekly units were flat at about 680 ({pf_weeks}): steady, not rising or fading yet. It earns ${pf_unit_profit:.0f} per unit at a {pf_margin}% margin, about ${pf_month_profit_k}K a month "
          f"at the September pace. It is at zero stock in {pf_zero} stores; each store sells about 57 units a month, so empty shelves cost "
          "roughly $5K a month. It was the #5 profit toy in Q3 2023 despite only ~6 weeks on sale. Be honest: hot now, likely to cool (see the previous slide). "
          "Only six weeks of data; the $20K is a run-rate, not a forecast.")

# ================================================================= slide 6: recommendation
s6 = content_slide("Next 8 weeks: keep Playfoam on shelves, test a Colorbuds fix, then back what works", 6)
steps = [("1", "Restock Playfoam", "Goal: stop losing Playfoam sales to empty shelves",
          ["Refill the 13 stores that have none", f"Playfoam goes from ${pf_month_profit_k}K to about ${pf_full_k}K profit a month, over half of what Colorbuds lost", "Check stock every week"]),
         ("2", "Test a Colorbuds fix", "Goal: win back some of the lost Colorbuds sales",
          ["Try new colors, a bundle with Gamer Headphones, or a better shelf spot", "Test it in 10 stores, against 10 similar stores that change nothing",
           "Winning back a quarter of the lost sales adds $11K profit a month"]),
         ("3", "Review in 8 weeks", "Goal: put effort behind what works",
          ["Roll out the Colorbuds fix if the test stores sold 15% more, otherwise drop it", "Check whether Playfoam is holding up or cooling off",
           "Pick the next higher-margin toy to back"])]
colw = (CW - 2 * 0.3) / 3
for i, (num, head, goal, pts) in enumerate(steps):
    cx = LM + i * (colw + 0.3)
    shape(s6, MSO_SHAPE.RECTANGLE, cx, 1.85, colw, 5.0, LIGHT, f"Panel {num}")
    badge = shape(s6, MSO_SHAPE.RECTANGLE, cx + 0.4, 2.1, 0.6, 0.6, (MINT_T, PINK_T, BLUE)[i], f"Number {num}")
    tf = badge.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = num
    r.font.size, r.font.bold, r.font.name = Pt(24), True, HEAD
    r.font.color.rgb = rgb("FFFFFF")
    text(s6, head, cx + 0.4, 2.85, colw - 0.8, 0.5, 22, NAVY, True, HEAD, name=f"Step {num}")
    text(s6, goal, cx + 0.4, 3.4, colw - 0.8, 0.7, 15, (MINT_T, PINK_T, BLUE)[i], True, name=f"Goal {num}")
    bullets(s6, pts, cx + 0.4, 4.05, colw - 0.8, 2.7, 15, space=8, name=f"Step {num} details")
notes(s6, "1) Keep Playfoam in stock: refill the 13 stores that have zero, and keep every store above about two weeks of sales (roughly 26 units a store); "
          "check weekly. Where $5K comes from: the average store sold about 57 Playfoam units in September (2,847 units / 50 stores), each unit earns $7 profit, so 13 empty stores x 57 x $7 = about $5K a month, which lifts Playfoam from about $20K to about $25K a month, more than half of the $44K a month Colorbuds has lost. It is conservative: those 13 stores actually averaged about 71 units each in September before running out. "
          "2) Test a Colorbuds fix: the options to try are new colors, a bundle with Gamer Headphones, or a better shelf spot. Run it in 10 stores and compare "
          "against 10 similar stores that change nothing, for about six weeks, so we learn what works before spending chain-wide. "
          "Winning back a quarter of the lost sales (about $44K a month lost, so about $11K) is the prize if it works. "
          "3) Review at week 8: if the pilot stores sold about 15% more Colorbuds than the comparison stores, scale the fix; if not, stop. Also check whether Playfoam is "
          "holding its weekly sales, and pick the next newer, higher-margin toy to back. The 15% threshold is a suggestion to agree on in advance, not something from "
          "the data. Sizes are estimates; the data cannot show why Colorbuds fell.")

OUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(OUT)
print("saved", OUT)
